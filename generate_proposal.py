# pip install python-docx
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_updated_proposal():
    doc = Document()

    # --- TITLE ---
    title = doc.add_heading('MarketX Precision Enterprise Suite', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('AI-Vision Diagnosis, Dynamic Cycle Tracking, Livestock Management, and Predictive Enterprise Planning for Kenyan Smallholders')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].bold = True

    # --- METADATA ---
    meta = doc.add_paragraph()
    meta.add_run('Target Grant Ask: ').bold = True
    meta.add_run('$45,000 USD\n')
    meta.add_run('Pilot Duration: ').bold = True
    meta.add_run('6 Months (1 Full Crop & Livestock Cycle)\n')
    meta.add_run('Target Beneficiaries: ').bold = True
    meta.add_run('500 Smallholder Farmers (~2,500 Household Beneficiaries)\n')
    meta.add_run('Applicant: ').bold = True
    meta.add_run('Kreek Analytics LLC (Operating MarketX / TerraPlan AI)')
    
    doc.add_paragraph('_' * 80)

    # --- SECTION 1: EXECUTIVE SUMMARY ---
    doc.add_heading('1. Executive Summary', level=1)
    doc.add_paragraph(
        "Smallholder farmers in Kenya face compounding risks: climate variability, extreme market opacity, and the catastrophic misuse of agrochemicals and veterinary drugs. "
        "Generic agricultural apps fail because they require high-bandwidth internet, ignore livestock, and lack offline physical fallbacks for rural realities."
    )
    doc.add_paragraph(
        "MarketX bridges this gap with a deeply integrated, offline-first agronomic and livestock intelligence suite. "
        "Built on a robust, pre-tested codebase, this 6-month pilot deploys core technical pillars directly to basic smartphones and feature phones:"
    )
    
    p1 = doc.add_paragraph(style='List Bullet')
    p1.add_run('Predictive Enterprise Planner (Crops & Livestock): ').bold = True
    p1.add_run('Generates multi-variable ROI models and printable, offline Farmer Handbooks for manual filing.')
    
    p2 = doc.add_paragraph(style='List Bullet')
    p2.add_run('Dynamic Cycle Ledger: ').bold = True
    p2.add_run('Weather-adjusted, day-by-day task calendars for both crop phenology and livestock husbandry.')
    
    p3 = doc.add_paragraph(style='List Bullet')
    p3.add_run('AI-Vision Diagnosis: ').bold = True
    p3.add_run('Edge-compressed image analysis and clinical decision trees for safe crop and animal disease treatment.')

    doc.add_paragraph(
        "This $45,000 grant will fund field deployment, Vision-AI compute, and econometric validation to prove a verifiable "
        "20%–30% net income gain and a 40% reduction in chemical/veterinary misuse for 500 farmers."
    )

    # --- SECTION 2: THE PROBLEM ---
    doc.add_heading('2. The Problem: The Agronomic & Livestock Triad of Failure', level=1)
    doc.add_paragraph("1. Blind Enterprise Planning: Farmers select crops or livestock based on tradition, lacking tools to model profitability, infrastructure costs, or local market forecasts.")
    doc.add_paragraph("2. Phenological & Husbandry Drift: Climate change and poor biosecurity cause farmers to miss critical weaning, vaccination, and scouting windows, leading to compounding losses.")
    doc.add_paragraph("3. Pathological Mismanagement: Farmers rely on 'chemical cocktails' for crops and unverified veterinary drugs for livestock, destroying soil biology, creating toxic residue, and violating Pre-Harvest/Withdrawal Intervals.")

    # --- SECTION 3: THE SOLUTION ---
    doc.add_heading('3. The MarketX Solution: 3 Core Technical Pillars', level=1)
    
    doc.add_heading('Pillar 1: The AI Enterprise Planner (Crops, Livestock & Offline Handbooks)', level=2)
    doc.add_paragraph(
        "Unlike generic advice, the MarketX Planner acts as a personalized farm economist for both agronomy and animal husbandry."
    )
    
    b1 = doc.add_paragraph(style='List Bullet')
    b1.add_run('Crops & The Printable Handbook: ').bold = True
    b1.add_run('The system models crop ROI based on GPS, soil, and budget. Crucially, it generates a comprehensive, printable "Farmer Handbook" (PDF) tailored to the specific crop. This allows farmers to keep a physical, manual filing copy for offline reference in the field, bridging the digital divide.')
    
    b2 = doc.add_paragraph(style='List Bullet')
    b2.add_run('Livestock Enterprise Module: ').bold = True
    b2.add_run('Expands beyond crops to dairy, poultry, and small ruminants. It provides exact infrastructure blueprints (e.g., zero-grazing unit dimensions), day-one arrival protocols (quarantine, stress-management, first vaccinations), full lifecycle vaccination cycles, and integrated disease prevention combining pharmaceutical schedules with non-pharmaceutical biosecurity measures.')

    doc.add_heading('Pillar 2: Dynamic Crop & Livestock Cycle Ledger', level=2)
    doc.add_paragraph(
        "Once an enterprise is selected, the system generates a dynamic, day-offset task calendar. "
        "For crops, it tracks land prep and fertilizer timing. For livestock, it tracks feed logs, milk yields, and vet visits. "
        "Every input is logged via an offline-first queue, creating a verifiable provenance ledger that prepares smallholders for premium B2B off-taker contracts and EUDR compliance."
    )

    doc.add_heading('Pillar 3: AI-Vision Diagnosis & Clinical Decision Tree', level=2)
    doc.add_paragraph("The most critical intervention for yield rescue and environmental/health safety.")
    
    b3 = doc.add_paragraph(style='List Bullet')
    b3.add_run('Low-Bandwidth Edge Compression: ').bold = True
    b3.add_run('Farmers snap a photo of a diseased leaf or a sick animal. The app uses client-side HTML5 canvas compression (max 1024px, JPEG 0.7) to ensure upload over rural 2G/3G networks.')
    
    b4 = doc.add_paragraph(style='List Bullet')
    b4.add_run('Vision AI & Differential Diagnosis: ').bold = True
    b4.add_run('The image is analyzed by Vision LLMs (e.g., Gemini Flash) to identify the pathology, severity, and affected organism.')
    
    b5 = doc.add_paragraph(style='List Bullet')
    b5.add_run('The Safety Decision Tree: ').bold = True
    b5.add_run('The AI output is passed through a strict clinical decision tree against a verified database of trusted agrochemicals and veterinary drugs. It outputs exact dosages, PPE requirements, and Pre-Harvest/Withdrawal Interval warnings. If uncertain, it defaults to safe biological/biosecurity advisories.')

    # --- SECTION 4: BUDGET ---
    doc.add_heading('4. Master $45,000 Itemized Budget', level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Budget Category'
    hdr_cells[1].text = 'Allocation ($)'
    hdr_cells[2].text = '%'
    hdr_cells[3].text = 'Technical Deliverables & Focus Area'
    
    for cell in hdr_cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True

    budget_data = [
        ('1. Ground Onboarding & Hardware', '$13,800', '30.7%', '• 10 Rugged Tablets & MDM ($2k)\n• Agent Stipends & Mobility ($9.2k)\n• Mobilization & Soil Kits ($2.6k)'),
        ('2. Cloud, AI Compute & Telco', '$10,200', '22.7%', '• Vision AI Inference Costs ($3.6k)\n• Offline-Sync Infrastructure ($3k)\n• USSD/SMS Fallbacks ($3.6k)'),
        ('3. Econometrics & M&E', '$9,500', '21.1%', '• Outcome Engine Variance Tracking ($4.5k)\n• Chemical/Vet Spend Audits ($3k)\n• ODK Survey Engineering ($2k)'),
        ('4. Capacity Building & Print Media', '$4,000', '8.9%', '• Printed Crop & Livestock Handbooks ($1.2k)\n• 2 Champion Demo Plots ($1.5k)\n• Field Days & Barazas ($1.3k)'),
        ('5. Verification & Case Study', '$2,500', '5.5%', '• Independent Agronomic/Vet Audit ($1.5k)\n• Public Impact Report ($1k)'),
        ('6. Legal & Admin Ops', '$5,000', '11.1%', '• County CECM permits, field micro-insurance, and financial compliance.'),
    ]

    for cat, amt, pct, deliv in budget_data:
        row_cells = table.add_row().cells
        row_cells[0].text = cat
        row_cells[1].text = amt
        row_cells[2].text = pct
        row_cells[3].text = deliv

    total_row = table.add_row().cells
    total_row[0].text = 'TOTAL MASTER BUDGET'
    total_row[1].text = '$45,000'
    total_row[2].text = '100%'
    total_row[3].text = '500 Farmers ($90 Unit Cost — 100% Free to Farmers)'
    for cell in total_row:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True

    # --- SECTION 5: KPIs ---
    doc.add_heading('5. Measurable Donor KPIs & B2B Transition', level=1)
    doc.add_paragraph("1. Economic Uplift: Statistically validated baseline vs. endline M&E data proving a 20%–30% net income increase per household across both crop and livestock enterprises.")
    doc.add_paragraph("2. Chemical & Vet Rationalization: A verifiable 40% reduction in unnecessary agro-chemical and veterinary spend, with zero withdrawal-interval violations.")
    doc.add_paragraph("3. Traceability Readiness: 100% of the 500 pilot farms will possess a digitized Enterprise Cycle Provenance Ledger.")
    doc.add_paragraph("4. Offline Adoption: >85% utilization rate of the printed Farmer Handbooks for manual record-keeping and offline reference.")

    doc.save('MarketX_Updated_Grant_Proposal.docx')
    print("Document created successfully: MarketX_Updated_Grant_Proposal.docx")

if __name__ == "__main__":
    create_updated_proposal()